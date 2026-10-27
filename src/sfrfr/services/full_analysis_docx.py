"""Сборка настоящего DOCX подробного разбора (этап 5) из Markdown-черновика."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt

_POSITION = (
    "Мы готовим документы и план — подаёте через СФР или Госуслуги вы сами. "
    "Решение принимает СФР."
)


def build_full_analysis_docx_bytes(
    *,
    case_ref: str,
    markdown: str,
    version: int = 1,
    title: str = "Подробный разбор стажа",
    note: str | None = None,
) -> bytes:
    """Реальный .docx (OOXML), не PDF с чужим расширением."""
    doc: Any = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading(title or "Подробный разбор стажа", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Дело (служебный ref): {case_ref or '—'} · Версия {int(version)}")
    note_p = doc.add_paragraph()
    note_p.add_run(
        note
        or (
            "Информационно-аналитический документ. Не является решением СФР. "
            "Клиент согласует текст; подача обращений — отдельный этап."
        )
    )

    body = (markdown or "").strip()
    if not body:
        doc.add_paragraph("(Текст разбора пуст — подготовьте черновик с ИИ и проверьте.)")
    else:
        _append_markdown(doc, body)

    lower = body.lower()
    if "решение принимает сфр" not in lower:
        doc.add_paragraph(_POSITION)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _append_markdown(doc: Any, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip() or " ", level=3)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip() or " ", level=2)
            continue
        if line.startswith("# "):
            # Уже есть главный заголовок — как подзаголовок
            doc.add_heading(line[2:].strip() or " ", level=2)
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            doc.add_paragraph(bullet.group(1), style="List Bullet")
            continue
        numbered = re.match(r"^(\d+)[.)]\s+(.+)$", line)
        if numbered:
            doc.add_paragraph(numbered.group(2), style="List Number")
            continue
        doc.add_paragraph(line)
