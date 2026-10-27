"""DOCX обращения из подтверждённых полей draft дела (без выдуманных фактов)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document

from sfrfr.services.appeals_plan import header_gate


def build_appeal_docx_bytes(
    *,
    title: str,
    body: str,
    case_ref: str | None = None,
    organ_name: str | None = None,
    organ_address: str | None = None,
    applicant_fio: str | None = None,
    representative: str | None = None,
    representation_confirmed: bool = False,
) -> bytes:
    """Один адресат на документ. Шапка — только проверенные поля; иначе не «к подаче»."""
    gate = header_gate(
        organ_name=organ_name,
        organ_address=organ_address,
        applicant_fio=applicant_fio,
        representative=representative,
        representation_confirmed=representation_confirmed,
    )
    doc = Document()
    heading = (title or "Проект обращения").strip() or "Проект обращения"
    doc.add_heading(heading, level=1)
    if case_ref:
        doc.add_paragraph(f"Дело (служебный ref): {case_ref}")
    if gate["organ_name"]:
        doc.add_paragraph(f"Адресат: {gate['organ_name']}")
    if gate["organ_address"]:
        doc.add_paragraph(f"Реквизиты / адрес: {gate['organ_address']}")
    if gate["applicant_fio"]:
        doc.add_paragraph(f"Заявитель: {gate['applicant_fio']}")
    if gate["representative"]:
        doc.add_paragraph(f"Представитель: {gate['representative']}")
    if not gate["ready_to_file"]:
        miss = ", ".join(gate["missing_header_fields"]) or "реквизиты"
        doc.add_paragraph(
            f"⚠ Не готов к подаче: заполните проверенные поля шапки ({miss}). "
            "Не выдумывать ФИО руководителя, адрес, доверенность или подпись."
        )
    doc.add_paragraph(
        "Документ подготовлен сервисом для самостоятельной подачи клиентом. "
        "Дата подготовки проекта не является датой подписи или подачи. "
        "Решение принимает СФР."
    )
    text = (body or "").strip()
    if not text:
        doc.add_paragraph("(Текст черновика пуст — заполните проект в карточке дела.)")
    else:
        for block in text.split("\n"):
            doc.add_paragraph(block if block.strip() else "")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def draft_payload_from_case(case: dict[str, Any]) -> tuple[str, str]:
    raw = case.get("draft")
    if isinstance(raw, dict):
        title = str(raw.get("title") or case.get("draft_title") or "Проект обращения")
        body = str(raw.get("body") or case.get("draft_body") or "")
    else:
        title = str(case.get("draft_title") or "Проект обращения")
        body = str(case.get("draft_body") or "")
    return title, body
