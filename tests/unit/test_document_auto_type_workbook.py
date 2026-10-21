"""Классификация трудовой: имя ТК, заголовок DOCX, таблицы."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from sfrfr.integrations.max.document_auto_type import guess_document_type, guess_from_bytes
from sfrfr.ocr.engine import extract_text_from_bytes


def _docx_bytes(
    *,
    paragraph: str | None = None,
    table_cell: str | None = None,
    title: str | None = None,
) -> bytes:
    doc = Document()
    if title:
        doc.core_properties.title = title
    if paragraph:
        doc.add_paragraph(paragraph)
    if table_cell:
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = table_cell
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_tk_filename_suggests_workbook() -> None:
    guess = guess_document_type(filename="ТК Лопакова Н.Ф. — копия.docx", text="")
    assert guess.doc_type == "workbook"
    assert guess.confidence >= 0.5


def test_docx_title_core_property_workbook() -> None:
    data = _docx_bytes(title="Трудовая книжка")
    text = extract_text_from_bytes(data, "scan.bin")  # суффикс врёт — sniff DOCX
    assert "трудовая" in text.lower()
    guess = guess_from_bytes(filename="scan.bin", data=data)
    assert guess.doc_type == "workbook"


def test_docx_table_header_workbook() -> None:
    data = _docx_bytes(table_cell="Трудовая книжка")
    guess = guess_from_bytes(filename="file.docx", data=data)
    assert guess.doc_type == "workbook"


def test_settle_treats_workbook_not_missing_labor() -> None:
    from sfrfr.integrations.max import document_background as bg

    text, complete = bg.build_final_message(
        [
            {"doc_type": "ils", "checksum_sha256": "1", "storage_path": "a.pdf"},
            {
                "doc_type": "workbook",
                "checksum_sha256": "2",
                "storage_path": "b.docx",
                "filename": "ТК копия.docx",
            },
        ]
    )
    assert complete is True
    assert "не хватает" not in text.lower()
