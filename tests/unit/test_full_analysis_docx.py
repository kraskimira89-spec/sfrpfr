"""DOCX подробного разбора — настоящий Office Open XML, не PDF."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from sfrfr.services.full_analysis_docx import build_full_analysis_docx_bytes


def test_docx_is_real_office_open_xml() -> None:
    raw = build_full_analysis_docx_bytes(
        case_ref="ПС-42",
        markdown=(
            "# Подробный разбор\n\n"
            "## Документы\n"
            "- ИЛС: ils.pdf\n\n"
            "## Расхождения\n"
            "1. Период 2004 не отражён.\n\n"
            "Решение принимает СФР.\n"
        ),
        version=1,
    )
    assert raw[:2] == b"PK"  # zip / OOXML
    doc = Document(BytesIO(raw))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("Подробный разбор" in t for t in texts)
    assert any("Период 2004" in t for t in texts)
    assert any("Решение принимает СФР" in t for t in texts)
    assert any("версия 1" in t.lower() or "Версия 1" in t for t in texts)


def test_empty_markdown_still_builds_docx() -> None:
    raw = build_full_analysis_docx_bytes(case_ref="x", markdown="", version=2)
    assert raw[:2] == b"PK"
    doc = Document(BytesIO(raw))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "пуст" in joined.lower() or "Пуст" in joined
